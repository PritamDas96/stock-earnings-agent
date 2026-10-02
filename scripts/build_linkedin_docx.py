"""Build a professional Word version of the LinkedIn post with screenshots.

Reads the post copy from docs/LINKEDIN_POST.txt and embeds the snapshot and
chart images from docs/assets, producing docs/LINKEDIN_POST.docx — a single
file ready to share or to attach images from.

Usage:
    python -m scripts.build_linkedin_docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

DOCS = Path(__file__).resolve().parent.parent / "docs"
ASSETS = DOCS / "assets"
POST = DOCS / "LINKEDIN_POST.txt"
OUT = DOCS / "LINKEDIN_POST.docx"

INK = RGBColor(0x11, 0x18, 0x27)
GREY = RGBColor(0x6B, 0x72, 0x80)
ACCENT = RGBColor(0x25, 0x63, 0xEB)

# (filename, caption) for the images attached beneath the copy.
IMAGES = [
    ("02_snapshot_output.png",
     "The full visual snapshot — metrics, gauges and six finance charts in one view."),
    ("11_chart_price_volume.png",
     "Price & volume — candlestick trend with a volume sub-panel."),
    ("12_chart_fundamentals_trend.png",
     "Revenue, net income & margins — is growth turning into profit?"),
    ("13_chart_analyst_targets.png",
     "Analyst price targets — consensus range and implied upside."),
    ("14_chart_valuation_multiples.png",
     "Valuation multiples — the standard lenses side by side."),
    ("15_chart_health_radar.png",
     "Financial-health profile — strengths and weaknesses in one shape."),
    ("10_chart_kpi_gauges.png",
     "KPI gauges — the three headline numbers, executive-style."),
]


def _set_base_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = INK


def main() -> None:
    doc = Document()
    _set_base_font(doc)

    title = doc.add_paragraph()
    run = title.add_run("Stock Earnings Agent — LinkedIn Post")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = INK

    sub = doc.add_paragraph()
    sub_run = sub.add_run("Ready-to-publish copy, followed by the images to attach")
    sub_run.italic = True
    sub_run.font.size = Pt(11)
    sub_run.font.color.rgb = GREY

    doc.add_paragraph()

    # Post copy, preserving blank lines as paragraph breaks.
    for line in POST.read_text(encoding="utf-8").splitlines():
        para = doc.add_paragraph()
        if line.strip():
            para.add_run(line)
        para.paragraph_format.space_after = Pt(6)

    # Images section.
    doc.add_page_break()
    heading = doc.add_paragraph()
    h_run = heading.add_run("Suggested images")
    h_run.bold = True
    h_run.font.size = Pt(15)
    h_run.font.color.rgb = ACCENT

    for name, caption in IMAGES:
        path = ASSETS / name
        if not path.exists():
            print(f"skip (missing): {name}")
            continue
        pic_para = doc.add_paragraph()
        pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic_para.add_run().add_picture(str(path), width=Inches(6.2))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = cap.add_run(caption)
        cap_run.italic = True
        cap_run.font.size = Pt(9.5)
        cap_run.font.color.rgb = GREY
        doc.add_paragraph()

    doc.save(str(OUT))
    print(f"saved {OUT} ({OUT.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
